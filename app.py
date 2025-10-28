import requests
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from PIL import Image, ImageDraw, ImageFont
from io import BytesIO
import json, os, re
from dotenv import load_dotenv

# Load environment variables
load_dotenv()
ZAI_API_KEY = os.getenv("ZAI_API_KEY")
UNSPLASH_ACCESS_KEY = os.getenv("UNSPLASH_ACCESS_KEY")

ZAI_API_URL = "https://api.z.ai/api/paas/v4/chat/completions"

# ==========================
# Professional Themes
# ==========================
THEMES = {
    "business": {
        "name": "Professional Blue",
        "bg": RGBColor(15, 32, 56),
        "title": RGBColor(255, 255, 255),
        "text": RGBColor(220, 230, 240),
        "accent": RGBColor(66, 133, 244),
        "footer_bg": RGBColor(10, 20, 35)
    },
    "tech": {
        "name": "Tech Dark",
        "bg": RGBColor(18, 18, 18),
        "title": RGBColor(0, 255, 200),
        "text": RGBColor(200, 200, 200),
        "accent": RGBColor(138, 43, 226),
        "footer_bg": RGBColor(10, 10, 10)
    },
    "creative": {
        "name": "Creative Vibrant",
        "bg": RGBColor(45, 26, 60),
        "title": RGBColor(255, 215, 0),
        "text": RGBColor(240, 240, 240),
        "accent": RGBColor(255, 105, 180),
        "footer_bg": RGBColor(30, 17, 40)
    },
    "nature": {
        "name": "Nature Green",
        "bg": RGBColor(25, 55, 45),
        "title": RGBColor(255, 255, 255),
        "text": RGBColor(220, 235, 225),
        "accent": RGBColor(76, 175, 80),
        "footer_bg": RGBColor(15, 35, 28)
    },
    "elegant": {
        "name": "Elegant Gray",
        "bg": RGBColor(50, 50, 55),
        "title": RGBColor(255, 255, 255),
        "text": RGBColor(220, 220, 220),
        "accent": RGBColor(192, 157, 102),
        "footer_bg": RGBColor(35, 35, 40)
    },
    "health": {
        "name": "Healthcare Blue",
        "bg": RGBColor(30, 50, 70),
        "title": RGBColor(255, 255, 255),
        "text": RGBColor(225, 235, 245),
        "accent": RGBColor(0, 168, 255),
        "footer_bg": RGBColor(20, 35, 50)
    }
}

def choose_theme(subject):
    """Auto-select theme based on presentation subject"""
    subject_lower = subject.lower()
    
    if any(word in subject_lower for word in ['business', 'finance', 'corporate', 'sales', 'marketing']):
        return THEMES['business']
    elif any(word in subject_lower for word in ['tech', 'ai', 'software', 'code', 'digital', 'cyber']):
        return THEMES['tech']
    elif any(word in subject_lower for word in ['creative', 'design', 'art', 'innovation', 'brand']):
        return THEMES['creative']
    elif any(word in subject_lower for word in ['nature', 'environment', 'eco', 'green', 'sustainability']):
        return THEMES['nature']
    elif any(word in subject_lower for word in ['health', 'medical', 'hospital', 'doctor', 'patient']):
        return THEMES['health']
    else:
        return THEMES['elegant']

# ==========================
# Image Functions
# ==========================
def get_unsplash_image(query):
    if not UNSPLASH_ACCESS_KEY:
        return None
    try:
        url = "https://api.unsplash.com/search/photos"
        headers = {"Authorization": f"Client-ID {UNSPLASH_ACCESS_KEY}"}
        params = {"query": query, "per_page": 1, "orientation": "landscape"}
        resp = requests.get(url, params=params, headers=headers, timeout=10).json()
        if resp.get("results"):
            img_url = resp["results"][0]["urls"]["regular"]
            print(f"  Downloading image for '{query}'...")
            return BytesIO(requests.get(img_url).content)
    except Exception as e:
        print(f"  Image fetch failed: {e}")
    return None

def get_placeholder_image(text="Image"):
    img = Image.new('RGB', (800, 600), color=(70, 130, 180))
    draw = ImageDraw.Draw(img)
    try: 
        font = ImageFont.truetype("arial.ttf", 40)
    except: 
        font = ImageFont.load_default()
    bbox = draw.textbbox((0,0), text, font=font)
    w,h = bbox[2]-bbox[0], bbox[3]-bbox[1]
    draw.text(((800-w)//2,(600-h)//2), text, fill=(255,255,255), font=font)
    bio = BytesIO()
    img.save(bio, format='PNG')
    bio.seek(0)
    return bio

# ==========================
# User input
# ==========================
print("=" * 50)
print("  PROFESSIONAL PRESENTATION GENERATOR")
print("=" * 50)

title = input("\nPresentation title: ")
description = input("Brief description: ")
presenter_name = input("Presenter name: ")
num_slides = int(input("Number of slides (e.g., 5): "))
use_images = input("Use images? (y/n): ").lower() == 'y'

# Auto-select theme
selected_theme = choose_theme(title + " " + description)
print(f"\n✓ Auto-selected theme: {selected_theme['name']}")

# ==========================
# AI Prompt for slides WITH BETTER NOTES
# ==========================
prompt = f"""
Generate {num_slides} professional slides about '{title}'. Description: {description}

For EACH slide, provide:
1. A clear, engaging title
2. 3-5 concise bullet points (keep each under 15 words)
3. An image keyword (one or two words describing what image would fit)
4. DETAILED speaker notes (2-3 sentences explaining what the presenter should say, include key talking points, examples, or statistics)

Return ONLY a JSON array like this:
[
  {{
    "title": "Introduction to AI",
    "content": "AI transforms industries\\nMachine learning enables predictions\\nDeep learning mimics human brain\\nApplications in healthcare and finance",
    "image_keyword": "artificial intelligence",
    "notes": "Welcome the audience and introduce the topic of artificial intelligence. Mention that AI has grown exponentially in the last decade, with applications ranging from healthcare diagnostics to financial forecasting. Share a brief statistic: the AI market is expected to reach $500 billion by 2024."
  }}
]

Make the notes conversational and helpful for the presenter. Include specific examples or data points they can mention.
"""

# ==========================
# Call Z.ai API
# ==========================
print("\nGenerating slides with AI...")
headers = {"Authorization": f"Bearer {ZAI_API_KEY}", "Content-Type": "application/json"}
body = {
    "model": "glm-4.6",
    "messages": [
        {"role": "system", "content": "You are an expert presentation designer. Create engaging content with detailed speaker notes."},
        {"role": "user", "content": prompt}
    ]
}

try:
    resp = requests.post(ZAI_API_URL, headers=headers, json=body, timeout=30)
    resp.raise_for_status()
    text_output = resp.json()["choices"][0]["message"]["content"]
    
    # Extract JSON
    json_match = re.search(r'\[.*\]', text_output, re.DOTALL)
    if not json_match:
        print("ERROR: No JSON found in response")
        exit(1)
    
    slides_data = json.loads(json_match.group())
    print(f"✓ Generated {len(slides_data)} slides with detailed notes")
    
except Exception as e:
    print(f"API Error: {e}")
    exit(1)

# ==========================
# Create Presentation
# ==========================
print("\nBuilding presentation...")
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

def add_footer(slide, theme, slide_number, presenter_name):
    """Add footer with slide number and presenter name"""
    # Footer background bar
    footer_bar = slide.shapes.add_shape(
        1,  # Rectangle
        0, 
        Inches(7.2), 
        prs.slide_width, 
        Inches(0.3)
    )
    footer_bar.fill.solid()
    footer_bar.fill.fore_color.rgb = theme["footer_bg"]
    footer_bar.line.fill.background()
    
    # Presenter name (left)
    name_box = slide.shapes.add_textbox(Inches(0.3), Inches(7.25), Inches(5), Inches(0.2))
    name_frame = name_box.text_frame
    name_frame.text = presenter_name
    name_para = name_frame.paragraphs[0]
    name_para.font.size = Pt(10)
    name_para.font.color.rgb = theme["text"]
    name_para.alignment = PP_ALIGN.LEFT
    
    # Slide number (right)
    num_box = slide.shapes.add_textbox(Inches(9), Inches(7.25), Inches(0.7), Inches(0.2))
    num_frame = num_box.text_frame
    num_frame.text = f"Slide {slide_number}"
    num_para = num_frame.paragraphs[0]
    num_para.font.size = Pt(10)
    num_para.font.color.rgb = theme["accent"]
    num_para.alignment = PP_ALIGN.RIGHT

# TITLE SLIDE
slide = prs.slides.add_slide(prs.slide_layouts[6])
background = slide.background
background.fill.solid()
background.fill.fore_color.rgb = selected_theme["bg"]

# Title with accent line
title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1.5))
title_frame = title_box.text_frame
title_frame.text = title.upper()
title_frame.paragraphs[0].font.size = Pt(48)
title_frame.paragraphs[0].font.bold = True
title_frame.paragraphs[0].font.color.rgb = selected_theme["title"]
title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# Accent line under title
accent = slide.shapes.add_shape(1, Inches(3.5), Inches(4.1), Inches(3), Inches(0.05))
accent.fill.solid()
accent.fill.fore_color.rgb = selected_theme["accent"]
accent.line.fill.background()

if description:
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(8), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = description
    subtitle_frame.paragraphs[0].font.size = Pt(20)
    subtitle_frame.paragraphs[0].font.color.rgb = selected_theme["text"]
    subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

# Presenter name on title slide
presenter_box = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(8), Inches(0.5))
presenter_frame = presenter_box.text_frame
presenter_frame.text = f"Presented by {presenter_name}"
presenter_frame.paragraphs[0].font.size = Pt(16)
presenter_frame.paragraphs[0].font.color.rgb = selected_theme["accent"]
presenter_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

print("✓ Title slide created")

# CONTENT SLIDES
for idx, slide_info in enumerate(slides_data, start=1):
    print(f"\nCreating slide {idx}...")
    
    # Create slide with background
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = selected_theme["bg"]
    
    # Title
    title_text = slide_info.get("title", f"Slide {idx}")
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.7))
    title_frame = title_box.text_frame
    title_frame.text = title_text
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(32)
    title_para.font.bold = True
    title_para.font.color.rgb = selected_theme["title"]
    print(f"  ✓ Title: {title_text}")
    
    # Accent line under title
    accent = slide.shapes.add_shape(1, Inches(0.5), Inches(1.15), Inches(1.5), Inches(0.04))
    accent.fill.solid()
    accent.fill.fore_color.rgb = selected_theme["accent"]
    accent.line.fill.background()
    
    # Content
    content_text = slide_info.get("content", "")
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.5), Inches(5))
    content_frame = content_box.text_frame
    content_frame.word_wrap = True
    
    content_lines = content_text.split('\n')
    for i, line in enumerate(content_lines):
        if line.strip():
            if i > 0:
                p = content_frame.add_paragraph()
            else:
                p = content_frame.paragraphs[0]
            p.text = "• " + line.strip()
            p.font.size = Pt(18)
            p.font.color.rgb = selected_theme["text"]
            p.space_before = Pt(10)
            p.space_after = Pt(6)
    print(f"  ✓ Content: {len(content_lines)} points")
    
    # Image
    if use_images:
        image_keyword = slide_info.get("image_keyword", "business")
        img = get_unsplash_image(image_keyword)
        if not img:
            img = get_placeholder_image(image_keyword)
        
        if img:
            try:
                slide.shapes.add_picture(img, Inches(5.5), Inches(1.5), width=Inches(4), height=Inches(4.5))
                print(f"  ✓ Image added")
            except Exception as e:
                print(f"  ✗ Image failed: {e}")
    
    # Enhanced Notes
    notes_text = slide_info.get("notes", "")
    if notes_text:
        # Add slide title to notes for context
        full_notes = f"SLIDE {idx}: {title_text}\n\n{notes_text}\n\n---\nKey Points:\n"
        for line in content_lines:
            if line.strip():
                full_notes += f"- {line.strip()}\n"
        
        slide.notes_slide.notes_text_frame.text = full_notes
        print(f"  ✓ Detailed notes added")
    
    # Footer with presenter name and slide number
    add_footer(slide, selected_theme, idx, presenter_name)


# Save

filename = f"presentation_{title.replace(' ', '_')[:30]}.pptx"
prs.save(filename)
print(f"\n{'='*50}")
print(f"✓✓✓ SUCCESS! Presentation created!")
print(f"{'='*50}")
print(f"File: {filename}")
print(f"Slides: {len(prs.slides)} (including title slide)")
print(f"Theme: {selected_theme['name']}")
print(f"Presenter: {presenter_name}")
print(f"Notes: Detailed speaker notes included")
print(f"{'='*50}")